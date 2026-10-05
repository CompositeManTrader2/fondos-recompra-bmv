"""
Reporte "BUYBACK ACTIVITY dd/mm/yyyy" con el formato de Punto Casa de Bolsa.

Columnas: TRADE DATE · STOCK · BROKER · B/S · SHARES · AVG PRICE · GROSS MXN
Una fila por (fecha de operación, emisora+serie, casa de bolsa, compra/venta),
ordenadas por emisora, fecha y lado — igual que el reporte de referencia.

Salidas: HTML (pantalla / correo) y Excel (.xlsx) con el mismo formato.
"""
from __future__ import annotations

import html
import io
from datetime import date
from pathlib import Path

import pandas as pd

PURPLE = "#7030A0"
ZEBRA = "#ECDEF5"
TEXTO = "#262626"
FUENTE = "Calibri, Carlito, 'Segoe UI', Arial, sans-serif"
LADO = {"COMPRA": "BOUGHT", "VENTA": "SOLD", "OTRO": "OTHER"}
COLUMNAS = ["TRADE DATE", "STOCK", "BROKER", "B/S", "SHARES", "AVG PRICE", "GROSS MXN"]
LOGO = Path(__file__).resolve().parent.parent / "assets" / "logo_punto.png"


def construir_tabla(actividad: pd.DataFrame, fecha: date, por: str = "reporte",
                    emisoras=None, casas=None, lados=None) -> pd.DataFrame:
    """
    por="reporte": documentos publicados ese día (como el reporte de Punto:
    el reporte del 11/03 incluye operaciones del 09 y 10/03).
    por="operacion": operaciones ejecutadas ese día.
    """
    if actividad is None or actividad.empty:
        return pd.DataFrame(columns=COLUMNAS)
    col = "FECHA_REPORTE" if por == "reporte" else "FECHA_OPERACION"
    d = actividad[pd.to_datetime(actividad[col]).dt.date == fecha].copy()
    if emisoras:
        d = d[d["EMISORA"].isin(emisoras)]
    if casas:
        d = d[d["CASA_BOLSA"].isin(casas)]
    if lados:
        d = d[d["TIPO"].isin(lados)]
    d = d[d["ACCIONES"].fillna(0) > 0]
    if d.empty:
        return pd.DataFrame(columns=COLUMNAS)
    d["_ORD"] = d["TIPO"].map({"COMPRA": 0, "VENTA": 1}).fillna(2)
    d = d.sort_values(["EMISORA", "FECHA_OPERACION", "_ORD", "CASA_BOLSA"])
    return pd.DataFrame({
        "TRADE DATE": pd.to_datetime(d["FECHA_OPERACION"]).dt.date.values,
        "STOCK": (d["EMISORA"] + " " + d["SERIE"].fillna("").astype(str)).str.strip().values,
        "BROKER": d["CASA_BOLSA"].values,
        "B/S": d["TIPO"].map(LADO).fillna(d["TIPO"]).values,
        "SHARES": d["ACCIONES"].astype(float).values,
        "AVG PRICE": d["PRECIO_PROM"].astype(float).values,
        "GROSS MXN": d["IMPORTE"].astype(float).values,
    })


# ---------------------------------------------------------------------------
# HTML
# ---------------------------------------------------------------------------

def _money(v: float, dec: int) -> str:
    """Formato contable: '$' a la izquierda de la celda, cifra a la derecha."""
    if v is None or v != v:
        return ""
    return (f'<span style="display:flex;justify-content:space-between;gap:10px">'
            f'<span>$</span><span>{v:,.{dec}f}</span></span>')


def html_tabla(tabla: pd.DataFrame, fecha: date, logo_html: str = "") -> str:
    th = (f"background:{PURPLE};color:#FFFFFF;font-weight:700;padding:3px 9px;border:none;"
          f"font-size:13px;white-space:nowrap;text-transform:uppercase")
    alin = {"TRADE DATE": "left", "STOCK": "left", "BROKER": "left", "B/S": "left",
            "SHARES": "right", "AVG PRICE": "center", "GROSS MXN": "center"}
    cab = "".join(f'<th style="{th};text-align:{alin[c]}">{c}</th>' for c in COLUMNAS)
    filas = []
    for i, r in enumerate(tabla.itertuples(index=False)):
        bg = "#FFFFFF" if i % 2 == 0 else ZEBRA
        td = f"padding:2px 9px;font-size:13px;color:{TEXTO};white-space:nowrap;background:{bg};border:none"
        filas.append(
            "<tr>"
            f'<td style="{td}">{r[0]:%d/%m/%Y}</td>'
            f'<td style="{td}">{html.escape(str(r[1]))}</td>'
            f'<td style="{td}">{html.escape(str(r[2]))}</td>'
            f'<td style="{td}">{html.escape(str(r[3]))}</td>'
            f'<td style="{td};text-align:right">{r[4]:,.0f}</td>'
            f'<td style="{td}">{_money(r[5], 4)}</td>'
            f'<td style="{td}">{_money(r[6], 2)}</td>'
            "</tr>"
        )
    cuerpo = "".join(filas) or (f'<tr><td colspan="7" style="padding:10px;font-size:13px;color:#6B6475">'
                                f'Sin operaciones de recompra para esta fecha.</td></tr>')
    return f"""
<div style="font-family:{FUENTE};display:inline-block;min-width:560px;max-width:100%;overflow-x:auto">
  <div style="display:flex;justify-content:space-between;align-items:flex-end;gap:24px;margin-bottom:6px">
    <div style="color:{PURPLE};font-weight:700;font-size:17px;letter-spacing:.02em;padding-bottom:4px">
      BUYBACK ACTIVITY {fecha:%d/%m/%Y}</div>
    <div>{logo_html}</div>
  </div>
  <table style="border-collapse:collapse;width:100%;border:none;margin:0">
    <thead><tr>{cab}</tr></thead><tbody>{cuerpo}</tbody>
  </table>
</div>"""


def html_documento(tabla: pd.DataFrame, fecha: date, logo_html: str = "") -> str:
    """Documento HTML autocontenido (para pegar en correo o abrir en navegador)."""
    return (f'<!doctype html><html><head><meta charset="utf-8"><title>Buyback Activity {fecha:%d/%m/%Y}</title>'
            f'</head><body style="margin:24px;background:#fff">{html_tabla(tabla, fecha, logo_html)}</body></html>')


# ---------------------------------------------------------------------------
# PNG (Pillow)
# ---------------------------------------------------------------------------

# Calibri (Windows) → Carlito (misma métrica que Calibri; apt fonts-crosextra-carlito
# en Streamlit Cloud) → DejaVu → fuente por defecto de Pillow.
_FUENTES = {
    "regular": ["C:/Windows/Fonts/calibri.ttf", "/usr/share/fonts/truetype/crosextra/Carlito-Regular.ttf",
                "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"],
    "bold": ["C:/Windows/Fonts/calibrib.ttf", "/usr/share/fonts/truetype/crosextra/Carlito-Bold.ttf",
             "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"],
}


def _fuente(peso: str, px: int):
    from PIL import ImageFont
    for ruta in _FUENTES[peso]:
        if Path(ruta).exists():
            return ImageFont.truetype(ruta, px)
    return ImageFont.load_default(size=px)


def _hex_rgb(h: str) -> tuple[int, int, int]:
    h = h.lstrip("#")
    return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))


def png_bytes(tabla: pd.DataFrame, fecha: date, escala: int = 2) -> bytes:
    """
    Imagen PNG de la tabla con el formato del reporte (título morado, logotipo,
    encabezado #7030A0, renglones alternos #ECDEF5, montos en formato contable).
    `escala=2` la genera a doble resolución para que se vea nítida en pantallas
    de alta densidad y al pegarla en correos o presentaciones.
    """
    from PIL import Image, ImageDraw

    s = escala
    f_txt, f_bold = _fuente("regular", 13 * s), _fuente("bold", 13 * s)
    f_tit = _fuente("bold", 17 * s)
    pad_x, pad_y, alto_fila, margen = 9 * s, 3 * s, 20 * s, 18 * s

    filas = [[f"{r[0]:%d/%m/%Y}", str(r[1]), str(r[2]), str(r[3]), f"{r[4]:,.0f}",
              f"{r[5]:,.4f}", f"{r[6]:,.2f}"] for r in tabla.itertuples(index=False)]
    medir = ImageDraw.Draw(Image.new("RGB", (1, 1)))
    ancho_txt = lambda t, f: medir.textlength(t, font=f)
    signo = ancho_txt("$", f_txt) + 10 * s          # "$" + separación en columnas de dinero
    anchos = []
    for j, col in enumerate(COLUMNAS):
        contenido = max((ancho_txt(f[j], f_txt) for f in filas), default=0)
        if j >= 5:
            contenido += signo
        anchos.append(int(max(ancho_txt(col, f_bold), contenido) + 2 * pad_x))
    ancho_tabla = sum(anchos)

    # Encabezado del reporte: título + logotipo
    alto_logo = 46 * s
    logo = None
    if LOGO.exists():
        try:
            logo = Image.open(LOGO).convert("RGBA")
            logo = logo.resize((int(logo.width * alto_logo / logo.height), alto_logo))
        except Exception:
            logo = None
    alto_cab = max(alto_logo, 30 * s) + 8 * s
    ancho = ancho_tabla + 2 * margen
    alto = margen + alto_cab + alto_fila * (len(filas) + 1) + margen + (24 * s if not filas else 0)
    img = Image.new("RGB", (int(ancho), int(alto)), "white")
    d = ImageDraw.Draw(img)
    morado, zebra, texto = _hex_rgb(PURPLE), _hex_rgb(ZEBRA), _hex_rgb(TEXTO)

    y_titulo = margen + alto_cab - 8 * s - (f_tit.size + 2 * s)
    d.text((margen, y_titulo), f"BUYBACK ACTIVITY {fecha:%d/%m/%Y}", font=f_tit, fill=morado)
    if logo is not None:
        img.paste(logo, (int(ancho - margen - logo.width), margen), logo)
    else:   # wordmark de texto: "punt" gris + "o" morada, "casa de bolsa" debajo
        gris = _hex_rgb("#949BA1")
        f_w, f_c = _fuente("bold", 30 * s), _fuente("regular", 12 * s)
        w_punt, w_o = ancho_txt("punt", f_w), ancho_txt("o", f_w)
        x0 = ancho - margen - (w_punt + w_o)
        d.text((x0, margen), "punt", font=f_w, fill=gris)
        d.text((x0 + w_punt, margen), "o", font=f_w, fill=morado)
        d.text((ancho - margen - ancho_txt("casa de bolsa", f_c), margen + 32 * s), "casa de bolsa", font=f_c, fill=gris)

    # Tabla
    y = margen + alto_cab
    d.rectangle([margen, y, margen + ancho_tabla, y + alto_fila], fill=morado)
    x = margen
    for j, col in enumerate(COLUMNAS):
        w = ancho_txt(col, f_bold)
        if col == "SHARES":
            tx = x + anchos[j] - pad_x - w
        elif j >= 5:
            tx = x + (anchos[j] - w) / 2
        else:
            tx = x + pad_x
        d.text((tx, y + pad_y), col, font=f_bold, fill="white")
        x += anchos[j]
    for i, f in enumerate(filas):
        y += alto_fila
        if i % 2 == 1:
            d.rectangle([margen, y, margen + ancho_tabla, y + alto_fila], fill=zebra)
        x = margen
        for j, valor in enumerate(f):
            w = ancho_txt(valor, f_txt)
            if j == 4:
                d.text((x + anchos[j] - pad_x - w, y + pad_y), valor, font=f_txt, fill=texto)
            elif j >= 5:      # formato contable: "$" a la izquierda, cifra a la derecha
                d.text((x + pad_x, y + pad_y), "$", font=f_txt, fill=texto)
                d.text((x + anchos[j] - pad_x - w, y + pad_y), valor, font=f_txt, fill=texto)
            else:
                d.text((x + pad_x, y + pad_y), valor, font=f_txt, fill=texto)
            x += anchos[j]
    if not filas:
        d.text((margen + pad_x, y + alto_fila + pad_y), "Sin operaciones de recompra para esta fecha.",
               font=f_txt, fill=_hex_rgb("#6B6475"))

    buf = io.BytesIO()
    img.save(buf, format="PNG", optimize=True, dpi=(96 * s, 96 * s))
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

def excel_bytes(tabla: pd.DataFrame, fecha: date) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = f"Buyback {fecha:%d-%m-%Y}"
    ws.sheet_view.showGridLines = False
    morado = PURPLE.lstrip("#")
    ws["A1"] = f"BUYBACK ACTIVITY {fecha:%d/%m/%Y}"
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color=morado)
    if LOGO.exists():
        try:
            from openpyxl.drawing.image import Image as XLImage
            img = XLImage(str(LOGO))
            escala = 45 / img.height
            img.height, img.width = 45, int(img.width * escala)
            ws.add_image(img, "E1")
        except Exception:
            pass

    fila_cab = 3
    relleno_cab = PatternFill("solid", fgColor=morado)
    relleno_zebra = PatternFill("solid", fgColor=ZEBRA.lstrip("#"))
    for j, c in enumerate(COLUMNAS, start=1):
        cell = ws.cell(row=fila_cab, column=j, value=c)
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.fill = relleno_cab
        cell.alignment = Alignment(horizontal="right" if c == "SHARES" else ("center" if j > 5 else "left"))

    formatos = {5: "#,##0", 6: '_-"$"* #,##0.0000_-', 7: '_-"$"* #,##0.00_-'}
    for i, r in enumerate(tabla.itertuples(index=False)):
        fila = fila_cab + 1 + i
        valores = [r[0], r[1], r[2], r[3], float(r[4]), float(r[5]), float(r[6])]
        for j, v in enumerate(valores, start=1):
            cell = ws.cell(row=fila, column=j, value=v)
            cell.font = Font(name="Calibri", size=11, color=TEXTO.lstrip("#"))
            if i % 2 == 1:
                cell.fill = relleno_zebra
            if j == 1:
                cell.number_format = "dd/mm/yyyy"
            elif j in formatos:
                cell.number_format = formatos[j]
    for j, ancho in enumerate([13, 13, 10, 10, 13, 13, 18], start=1):
        ws.column_dimensions[chr(64 + j)].width = ancho
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
